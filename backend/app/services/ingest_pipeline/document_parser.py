import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path


MAX_DOCUMENT_TEXT_CHARS = 1_000_000
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class DocumentExtractionError(ValueError):
    pass


class DocumentParserUnavailableError(RuntimeError):
    pass


def _limit_text(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_DOCUMENT_TEXT_CHARS:
        raise DocumentExtractionError(
            f"Extracted document text exceeds {MAX_DOCUMENT_TEXT_CHARS:,} characters."
        )
    if not text:
        raise DocumentExtractionError("No readable text was found in the document.")
    return text


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=False)
    except ImportError as exc:
        raise DocumentParserUnavailableError(
            "PDF support is not installed. Install the backend requirements and restart the service."
        ) from exc

    try:
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentExtractionError("Password-protected PDFs are not supported.")
        if len(reader.pages) > 2_000:
            raise DocumentExtractionError("PDFs may contain no more than 2,000 pages.")

        pages = []
        character_count = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            character_count += len(page_text)
            if character_count > MAX_DOCUMENT_TEXT_CHARS:
                raise DocumentExtractionError(
                    f"Extracted document text exceeds {MAX_DOCUMENT_TEXT_CHARS:,} characters."
                )
            pages.append(page_text)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError("Could not read this PDF file.") from exc
    return _limit_text("\n\n".join(pages))


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 5_000
                or sum(info.file_size for info in entries)
                > MAX_DOCX_UNCOMPRESSED_BYTES
            ):
                raise DocumentExtractionError("The DOCX file expands beyond the supported size.")
        document = Document(io.BytesIO(data))
    except (DocumentExtractionError, DocumentParserUnavailableError):
        raise
    except ImportError as exc:
        raise DocumentParserUnavailableError(
            "DOCX support is not installed. Install the backend requirements and restart the service."
        ) from exc
    except Exception as exc:
        raise DocumentExtractionError("Could not read this DOCX file.") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return _limit_text("\n\n".join(blocks))


def _extract_doc(data: bytes) -> str:
    antiword = shutil.which("antiword")
    if antiword is None:
        raise DocumentExtractionError(
            "Legacy .doc files require antiword; upload a PDF or DOCX instead."
        )

    try:
        with tempfile.NamedTemporaryFile(suffix=".doc") as document_file:
            document_file.write(data)
            document_file.flush()
            result = subprocess.run(
                [antiword, document_file.name],
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=30,
            )
    except subprocess.TimeoutExpired as exc:
        raise DocumentExtractionError("Reading this DOC file timed out.") from exc
    except OSError as exc:
        raise DocumentExtractionError("Could not read this DOC file.") from exc

    if result.returncode != 0:
        raise DocumentExtractionError("Could not read this DOC file.")
    return _limit_text(result.stdout)


def extract_document_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".doc":
        return _extract_doc(data)
    raise DocumentExtractionError("Only PDF, DOC, and DOCX files are supported.")
