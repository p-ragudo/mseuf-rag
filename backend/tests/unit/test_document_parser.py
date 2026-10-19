import io
import builtins

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.ingest_pipeline.document_parser import (
    DocumentExtractionError,
    DocumentParserUnavailableError,
    extract_document_text,
)


def test_extracts_docx_paragraphs_and_table_cells():
    document = Document()
    document.add_paragraph("Admissions are open for the upcoming academic year.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Application form"
    table.cell(0, 1).text = "Submit to the registrar"
    output = io.BytesIO()
    document.save(output)

    text = extract_document_text("admissions.docx", output.getvalue())

    assert "Admissions are open for the upcoming academic year." in text
    assert "Application form | Submit to the registrar" in text


def test_rejects_pdf_without_extractable_text():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    output = io.BytesIO()
    writer.write(output)

    with pytest.raises(DocumentExtractionError, match="No readable text"):
        extract_document_text("empty.pdf", output.getvalue())


def test_rejects_unsupported_file_type():
    with pytest.raises(DocumentExtractionError, match="Only PDF, DOC, and DOCX"):
        extract_document_text("notes.txt", b"plain text")


@pytest.mark.parametrize(
    ("filename", "missing_module"),
    [("admissions.pdf", "pypdf"), ("admissions.docx", "docx")],
)
def test_missing_parser_dependency_is_reported_at_upload_time(
    monkeypatch, filename, missing_module
):
    original_import = builtins.__import__

    def import_without_parser(name, *args, **kwargs):
        if name == missing_module:
            raise ModuleNotFoundError(f"No module named '{missing_module}'")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", import_without_parser)

    with pytest.raises(DocumentParserUnavailableError, match="support is not installed"):
        extract_document_text(filename, b"not parsed")
